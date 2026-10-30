"""MANUSCRIPT_BMC_v8.md -> submission .docx by the BMC Bioinformatics rules for the main manuscript text:
double line spacing, continuous line numbers, page numbers, no page breaks, tables as Word tables without shading,
Unicode characters (no symbol fonts). [AUTHOR ...] placeholders are highlighted in yellow.
Usage: python docs/make_submission_docx.py docs/MANUSCRIPT_BMC_v8.md out.docx   (needs python-docx)"""
import re
import sys

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

src, out = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().split("\n")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.line_spacing = 2.0
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)

sec = doc.sections[0]
ln = OxmlElement("w:lnNumType")                       # continuous line numbering
ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "360")
sec._sectPr.append(ln)
fp = sec.footer.paragraphs[0]                          # page number field
fp.alignment = 1
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if kind:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text; r._r.append(e)

TOKEN = re.compile(r"(\[AUTHOR[^\]]*\]|\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_runs(p, text, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("[AUTHOR"):
            r = p.add_run(part); r.font.highlight_color = WD_COLOR_INDEX.YELLOW; r.bold = bold
        elif part.startswith("**"):
            add_runs(p, part[2:-2], bold=True, italic=italic)
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Courier New"; r.bold = bold; r.italic = italic
        elif part.startswith("*") and len(part) > 2:
            add_runs(p, part[1:-1], bold=bold, italic=True)
        else:
            r = p.add_run(part); r.bold = bold; r.italic = italic


def heading(text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True
    r.font.size = Pt({1: 14, 2: 12, 3: 12}[level]); r.italic = level == 3


def table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|[-| ]+\|$", r.strip())]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, c in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, c, bold=(i == 0))


i = 0
title_done = False
while i < len(lines):
    l = lines[i]
    if l.startswith("# "):                             # title of the document = the title-page title
        i += 1; continue
    if l == "## Title page":
        i += 1; continue
    if l.startswith("**Title.** "):
        p = doc.add_paragraph(); r = p.add_run(l[len("**Title.** "):]); r.bold = True; r.font.size = Pt(16)
        title_done = True; i += 1; continue
    if l.startswith("|"):
        block = []
        while i < len(lines) and lines[i].startswith("|"):
            block.append(lines[i]); i += 1
        table(block); continue
    if l.startswith("### "):
        heading(l[4:], 3)
    elif l.startswith("## "):
        heading(l[3:], 2)
    elif l.startswith("- "):
        p = doc.add_paragraph(); add_runs(p, l[2:])
    elif l.strip():
        p = doc.add_paragraph(); add_runs(p, l)
    i += 1

assert title_done
doc.save(out)
print("wrote", out)
